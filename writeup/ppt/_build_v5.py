"""Build manuscript V5 from V4: text fixes + remove done annotations + swap redrawn figures."""
import io, shutil, re
from pathlib import Path
import docx
from docx.oxml.ns import qn
from PIL import Image

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
SRC = HERE / 'CAP_sleep_mask_manuscript_V4.docx'
DST = HERE / 'CAP_sleep_mask_manuscript_V5.docx'
shutil.copyfile(SRC, DST)
d = docx.Document(DST)
P = d.paragraphs

# ── 1. Text fixes ────────────────────────────────────────────────────────────
# para 85: (Figure S4) -> (Figure S12)   [asymmetry integral = imbalance figure]
r = P[85].runs[8]; assert r.text == '4', r.text; r.text = '12'
# para 194: estimate (Figure S3) -> S2   [ten-minute calibration = calibration-strategies fig]
r = P[194].runs[18]; assert 'Figure S3' in r.text; r.text = r.text.replace('Figure S3', 'Figure S2')
# para 191: insert the missing calibration factor k (italic, matching §4.4 usage)
def insert_italic_k_before(run, keep_leading_space=True):
    """Split so an italic 'k ' is inserted at the START of `run` text 'of ...'."""
    # run text begins 'of ...'; we want 'k of ...' with k italic
    new = run._r.makeelement(qn('w:r'), {})
    # copy rPr then set italic
    rpr = run._r.find(qn('w:rPr'))
    import copy
    if rpr is not None:
        new.append(copy.deepcopy(rpr))
    it = new.find(qn('w:rPr'))
    if it is None:
        it = new.makeelement(qn('w:rPr'), {}); new.insert(0, it)
    it.append(new.makeelement(qn('w:i'), {}))
    t = new.makeelement(qn('w:t'), {qn('xml:space'): 'preserve'}); t.text = 'k'
    new.append(t)
    run._r.addprevious(new)
    # add a plain space run 'k '<space> -> actually put space as normal run before 'of'
    sp = run._r.makeelement(qn('w:r'), {})
    st = sp.makeelement(qn('w:t'), {qn('xml:space'): 'preserve'}); st.text = ' '
    sp.append(st)
    run._r.addprevious(sp)

# Insert into the LATER run first so earlier-run indices don't shift.
# r2 = 'of 1.18 ...' (after 'the respiratory '); r1 = 'of approximately 2 ...' (after 'The cardiac ')
assert P[191].runs[1].text.startswith('of approximately 2')
assert P[191].runs[2].text.startswith('of 1.18')
insert_italic_k_before(P[191].runs[2])   # respiratory k
insert_italic_k_before(P[191].runs[1])   # cardiac k

# ── 2. Remove completed editorial annotations (redraw/font notes) ────────────
# In each figure-caption paragraph the ONLY highlighted content is the annotation
# ("(LARGER font size)" etc.), which can be split across runs — so empty every
# highlighted run in those paragraphs.
from docx.enum.text import WD_COLOR_INDEX as _WC
ANN_PARAS = (90, 110, 127, 130, 140, 151, 158, 163)
n_ann = 0
for pi in ANN_PARAS:
    for run in P[pi].runs:
        hc = run.font.highlight_color
        if hc is not None and hc != _WC.AUTO:
            if run.text.strip():
                n_ann += 1
            run.text = ''
    # drop a dangling trailing space left where the annotation was
    if P[pi].runs and P[pi].runs[-1].text == ' ':
        P[pi].runs[-1].text = ''

# ── 3. Clear highlight on the runs we resolved ───────────────────────────────
from docx.enum.text import WD_COLOR_INDEX
def clear_hl(run):
    run.font.highlight_color = None
for ri in (7, 8, 9):        # para 85 (Figure S12)
    clear_hl(P[85].runs[ri])
for run in P[191].runs:     # para 191 whole (k inserted)
    clear_hl(run)
clear_hl(P[194].runs[18])   # para 194 (Figure S2)
# annotation paragraphs: clear any now-empty highlighted runs
for pi in (90, 110, 127, 130, 140, 151, 158, 163):
    for run in P[pi].runs:
        if run.text.strip() == '':
            clear_hl(run)

# ── 4. Swap redrawn figures (blob + aspect-preserving extent) ────────────────
FIG = ROOT / 'writeup' / 'figures'
SWAP = {
    'rId11': FIG / 'signal_validation' / 'fig2_inband_snr.png',        # Fig 3
    'rId12': FIG / 'rate_rerun' / 'fig_representative_night.png',       # Fig 4
    'rId14': FIG / 'harmonics' / 'ridges' / 'ridge_tune_S1N2_CRE.png',  # Fig 5
    'rId15': FIG / 'harmonics' / 'band_ridge_by_stage.png',            # Fig 6
    'rId16': FIG / 'spindles' / 'fig_spindle_lowband_allsessions.png',  # Fig 7
    'rId17': FIG / 'delta_onset' / 'fig_delta_onset_cohort.png',        # Fig 8
    'rId18': FIG / 'harmonics' / 'ladders' / 'ladder_S6N1.png',         # Fig 9
    'rId19': FIG / 'harmonics' / 'ladder_stage_relationship.png',       # Fig 10
}
body = d.element.body
for rid, png in SWAP.items():
    assert png.exists(), f'missing {png}'
    newbytes = png.read_bytes()
    part = d.part.related_parts[rid]
    part._blob = newbytes
    w, h = Image.open(io.BytesIO(newbytes)).size
    aspect = h / w
    # find the drawing that embeds this rId and fix its extents (keep width, set height)
    for blip in body.iter(qn('a:blip')):
        if blip.get(qn('r:embed')) != rid:
            continue
        # ascend to the inline/anchor drawing container
        anc = blip
        while anc is not None and anc.tag not in (qn('wp:inline'), qn('wp:anchor')):
            anc = anc.getparent()
        if anc is None:
            continue
        ext = anc.find(qn('wp:extent'))
        if ext is not None:
            cx = int(ext.get('cx')); ext.set('cy', str(int(round(cx * aspect))))
        for aext in anc.iter(qn('a:ext')):
            if aext.get('cx') is not None:
                cx = int(aext.get('cx')); aext.set('cy', str(int(round(cx * aspect))))
        break

d.save(DST)
print('saved', DST.name)
print(f'annotations removed: {n_ann}')
# verify
d2 = docx.Document(DST)
print('para85:', d2.paragraphs[85].text[130:180])
print('para191 head:', d2.paragraphs[191].text[:70])
print('para194 has S2:', 'Figure S2' in d2.paragraphs[194].text, '| has S3:', '(Figure S3)' in d2.paragraphs[194].text)
print('para85 has S12:', 'Figure S12' in d2.paragraphs[85].text)
print('para90 tail:', repr(d2.paragraphs[90].text[-40:]))
print('images:', sum(1 for r in d2.part.rels.values() if 'image' in r.reltype))
# any leftover annotations?
import re as _re
for pi,p in enumerate(d2.paragraphs):
    if _re.search(r'LARGER font|Larger font|Redraw this', p.text):
        print('  LEFTOVER annotation in para', pi, ':', repr(p.text[-60:]))
